"""生成入库测试用的 txt/docx/pdf/md 小样本（不依赖网络与真实模型）。

用法::

    python tests/fixtures/make_samples.py <输出目录>

测试里直接 import 本模块的 ``build_docx`` / ``build_pdf`` / ``build_md`` 使用临时目录。
"""
from __future__ import annotations

import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
TXT_SAMPLE = HERE / "sample_policy.txt"


def build_txt(dst_dir: Path) -> Path:
    """复制一份纯文本制度样例。"""
    dst_dir = Path(dst_dir)
    dst_dir.mkdir(parents=True, exist_ok=True)
    target = dst_dir / "差旅费管理办法.txt"
    target.write_text(TXT_SAMPLE.read_text(encoding="utf-8"), encoding="utf-8")
    return target


def build_docx(dst_dir: Path) -> Path:
    """生成带章节、条款和 Word 表格的 DOCX。"""
    from docx import Document

    dst_dir = Path(dst_dir)
    dst_dir.mkdir(parents=True, exist_ok=True)
    target = dst_dir / "科研经费报销细则.docx"

    document = Document()
    document.add_paragraph("XX大学科研经费报销实施细则")
    document.add_paragraph("")
    document.add_paragraph("第一章 总则")
    document.add_paragraph("第一条 为规范科研经费报销，制定本细则。")
    document.add_paragraph("第二条 本细则适用于全校科研项目经费支出。")
    document.add_paragraph("第二章 报销标准")
    document.add_paragraph("第三条 科研经费支出按照以下标准执行：")
    document.add_paragraph("科研经费报销标准表")
    table = document.add_table(rows=3, cols=3)
    rows = [
        ("项目", "标准", "说明"),
        ("差旅费", "300元/天", "凭票据报销"),
        ("会议费", "500元/人", "需附会议通知"),
    ]
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value
    document.add_paragraph("第四条 表格所列标准由财务处负责解释。")
    document.add_paragraph("附件1：科研经费报销审批单")
    document.add_paragraph("报销审批单应当经项目负责人签字后提交财务处。")
    document.save(str(target))
    return target


def build_pdf(dst_dir: Path) -> Path:
    """生成两页文字版 PDF，用于验证页码定位。"""
    import fitz

    dst_dir = Path(dst_dir)
    dst_dir.mkdir(parents=True, exist_ok=True)
    target = dst_dir / "实验室安全管理办法.pdf"

    document = fitz.open()
    font = fitz.Font("cjk")

    page1 = document.new_page()
    page1.insert_font(fontname="cjk", fontbuffer=font.buffer)
    lines1 = [
        "XX大学实验室安全管理办法",
        "",
        "第一章 总则",
        "第一条 为加强实验室安全管理，保障师生人身安全，制定本办法。",
        "第二条 本办法适用于学校各类实验室及实验场所。",
    ]
    y = 72
    for line in lines1:
        if line:
            page1.insert_text((72, y), line, fontname="cjk", fontsize=12)
        y += 24
    page1.insert_text((240, 780), "第 1 页 / 共 2 页", fontname="cjk", fontsize=9)

    page2 = document.new_page()
    page2.insert_font(fontname="cjk", fontbuffer=font.buffer)
    lines2 = [
        "第二章 危险化学品管理",
        "第三条 危险化学品的领用应当经实验室负责人审批，单次领用不超过 5 天用量。",
        "第四条 危险化学品应当分类存放，专人保管，并建立台账。",
    ]
    y = 72
    for line in lines2:
        page2.insert_text((72, y), line, fontname="cjk", fontsize=12)
        y += 24
    page2.insert_text((240, 780), "第 2 页 / 共 2 页", fontname="cjk", fontsize=9)

    document.save(str(target))
    document.close()
    return target


def build_md(dst_dir: Path) -> Path:
    """生成带 Markdown 标题与管道表格的 md 样例。"""
    dst_dir = Path(dst_dir)
    dst_dir.mkdir(parents=True, exist_ok=True)
    target = dst_dir / "请假与考勤规定.md"
    target.write_text(
        "# XX大学教职工请假与考勤管理规定\n\n"
        "## 第一章 总则\n\n"
        "第一条 为规范教职工请假与考勤管理，制定本规定。\n\n"
        "## 第二章 审批权限\n\n"
        "第二条 请假审批权限如下：\n\n"
        "| 请假天数 | 审批人 | 备注 |\n"
        "| --- | --- | --- |\n"
        "| 1天以内 | 科室负责人 | 口头报备 |\n"
        "| 1至3天 | 部门负责人 | 书面申请 |\n"
        "| 3天以上 | 分管校领导 | 人事处备案 |\n\n"
        "第三条 未经批准擅自离岗的，按旷工处理。\n",
        encoding="utf-8",
    )
    return target


def build_all(dst_dir: Path) -> dict[str, Path]:
    return {
        "txt": build_txt(dst_dir),
        "docx": build_docx(dst_dir),
        "pdf": build_pdf(dst_dir),
        "md": build_md(dst_dir),
    }


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "generated"
    for kind, path in build_all(target).items():
        print(f"{kind}: {path}")
