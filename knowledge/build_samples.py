# -*- coding: utf-8 -*-
"""生成 knowledge/samples 下的 5 份中文制度样例及 metadata.json。

用法（在项目根目录）：
    .venv\\Scripts\\python.exe knowledge\\build_samples.py

依赖：PyMuPDF（fitz）、python-docx。PDF 使用 Windows 自带微软雅黑，
确保 fitz 可正常提取中文；DOCX 使用 python-docx 生成，含表格与附件。
生成后 metadata.json 中的 doc_id 按契约 doc-<sha256前12位> 计算。
"""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

import fitz
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SAMPLES = ROOT / "samples"
MSYH = r"C:\Windows\Fonts\msyh.ttc"
MSYH_BD = r"C:\Windows\Fonts\msyhbd.ttc"

PAGE_W, PAGE_H = 595, 842
MARGIN_X = 72
BODY_TOP = 96
BODY_BOTTOM = 780


# --------------------------------------------------------------------------
# 文档内容：统一用 (kind, text) 序列，kind = title/meta/chapter/article/note
# --------------------------------------------------------------------------
DOC1 = {
    "file": "XX大学实验室安全管理办法（2023年版）.pdf",
    "title": "XX大学实验室安全管理办法",
    "doc_no": "校发〔2023〕12号",
    "version_label": "2023年版",
    "effective_from": "2023-03-01",
    "effective_to": "2025-08-31",
    "department": "实验室与设备管理处",
    "status": "superseded",
    "blocks": [
        ("chapter", "第一章  总则"),
        ("article", "第一条  为加强学校实验室安全管理，预防和减少实验室安全事故，保障师生员工人身安全和学校财产安全，根据《中华人民共和国安全生产法》《危险化学品安全管理条例》《高等学校实验室工作规程》等法律法规和上级有关规定，结合学校实际，制定本办法。"),
        ("article", "第二条  本办法适用于学校各学院、研究院（所）、工程训练中心及直属单位设立的教学、科研实验室，以及经学校批准在校内设立的联合实验室、公共实验平台等场所。"),
        ("article", "第三条  实验室安全工作坚持“安全第一、预防为主、综合治理”的方针，落实“党政同责、一岗双责、齐抓共管、失职追责”的要求，实行学校统一领导、职能部门监督管理、学院直接管理、实验室具体负责、师生全员参与的工作机制。"),
        ("chapter", "第二章  安全管理责任"),
        ("article", "第四条  学校成立实验室安全工作领导小组，由分管实验室工作的校领导任组长，实验室与设备管理处、保卫处、教务处、科学技术研究院、研究生院、后勤保障处等部门负责人为成员。领导小组负责审议实验室安全规划与制度，统筹协调重大事项，组织开展安全检查与事故调查。"),
        ("article", "第五条  学院主要负责人是本单位实验室安全第一责任人，分管实验室工作的负责人是直接责任人；实验室负责人对本实验室的安全管理负具体责任。各实验室应当配备专（兼）职安全员，协助开展日常检查、台账登记和隐患整改。"),
        ("article", "第六条  实验室与设备管理处负责实验室安全工作的归口管理，承担危险化学品采购与储存监管、安全培训组织、检查考核和事故调查等工作；保卫处负责消防、治安及突发事件应急处置；教务处、科学技术研究院、研究生院按照职责分工，负责教学、科研和研究生实验活动的安全审核与督促。"),
        ("chapter", "第三章  危险化学品管理"),
        ("article", "第七条  危险化学品实行统一采购、集中储存、分级管理。任何单位和个人不得私自采购、接收、转让、出借危险化学品，不得通过非正规渠道购买易制毒、易制爆化学品。"),
        ("article", "第八条  危险化学品应当储存在符合国家标准的专用库房或者储存柜内，做到分类分区、专柜存放、双人双锁，并建立出入库台账，如实记录品名、规格、数量、领用人和领用时间，做到账物相符。"),
        ("article", "第九条  危险化学品领用应当经实验室负责人批准，由两人共同领取并签字登记，单次领用量不得超过5个工作日用量。使用剩余的危险化学品应当及时退回仓库，不得在实验室超量存放；严禁将危险化学品带出实验室或者转交无关人员。"),
        ("article", "第十条  危险废物应当分类收集、定点存放，粘贴规范标签并如实填写名称、成分、危险特性和产生日期，由学校委托具有相应资质的单位统一处置。严禁随意倾倒、填埋、焚烧或者排入下水道。"),
        ("chapter", "第四章  安全培训与检查"),
        ("article", "第十一条  新进入实验室的教师、学生和工作人员必须接受不少于8学时的实验室安全培训，经考试合格后方可进入实验室开展实验；在岗人员每年应当接受不少于2学时的安全再教育。涉及危险化学品、辐射、生物等特殊领域的，还应当接受专项培训。"),
        ("article", "第十二条  实验室实行安全自查制度。实验室每周自查一次，学院每月检查一次，学校每季度组织抽查。检查应当形成记录，发现隐患立即整改；不能立即整改的，应当明确责任人、整改措施和完成期限，并采取临时安全防护措施。"),
        ("article", "第十三条  发生实验室安全事故，现场人员应当立即采取应急处置措施，防止事态扩大，并在2小时内报告实验室与设备管理处和保卫处，不得迟报、漏报、瞒报。发生人员伤亡或者重大财产损失的，按照学校突发事件应急预案处置。"),
        ("chapter", "第五章  附则"),
        ("article", "第十四条  本办法由实验室与设备管理处负责解释。"),
        ("article", "第十五条  本办法自2023年3月1日起施行，原《XX大学实验室安全管理办法（试行）》同时废止。"),
    ],
}

DOC2 = {
    "file": "XX大学实验室安全管理办法（2025年修订）.docx",
    "title": "XX大学实验室安全管理办法（2025年修订）",
    "doc_no": "校发〔2025〕27号",
    "version_label": "2025年修订版",
    "effective_from": "2025-09-01",
    "effective_to": None,
    "department": "实验室与设备管理处",
    "status": "active",
    "blocks": [
        ("chapter", "第一章  总则"),
        ("article", "第一条  为加强学校实验室安全管理，预防和减少实验室安全事故，保障师生员工人身安全和学校财产安全，根据《中华人民共和国安全生产法》《危险化学品安全管理条例》《高等学校实验室安全规范》等法律法规和上级有关规定，结合学校实际，制定本办法。"),
        ("article", "第二条  本办法适用于学校各学院、研究院（所）、工程训练中心、公共实验平台、校外研究院及实习实训基地设立的教学科研实验室及其附属场所。"),
        ("article", "第三条  实验室安全工作坚持安全第一、预防为主、综合治理的方针，落实全员安全生产责任制，实行学校统一领导、职能部门监督管理、学院直接管理、实验室具体负责、师生全员参与的工作机制。"),
        ("chapter", "第二章  安全管理责任"),
        ("article", "第四条  学校实验室安全工作领导小组由校长任组长，分管实验室工作的校领导任副组长，实验室与设备管理处、保卫处、教务处、科学技术研究院、研究生院、学生工作部、人事处、后勤保障处等部门负责人为成员，负责统筹规划、审议制度、组织检查和事故调查。"),
        ("article", "第五条  学院主要负责人是本单位实验室安全第一责任人，分管负责人是直接责任人；实验室负责人对本实验室安全负具体责任。各实验室应当配备专（兼）职安全员，安全员协助开展日常检查、台账登记和隐患整改。"),
        ("article", "第六条  实验室与设备管理处负责实验室安全归口管理，承担危险化学品采购与储存监管、安全培训组织、检查考核和事故调查等工作；保卫处负责消防、治安和应急处置；教务处、科学技术研究院、研究生院负责教学科研活动的安全审核；人事处将实验室安全培训纳入教职工年度考核。"),
        ("chapter", "第三章  危险化学品与危险废物"),
        ("article", "第七条  危险化学品实行线上统一采购、集中储存、分级管理。任何单位和个人不得私自采购、接收、转让、出借危险化学品；易制毒、易制爆化学品应当按规定向公安机关备案。"),
        ("article", "第八条  危险化学品应当储存在符合国家标准的专用库房或者储存柜内，做到分类分区、专柜存放、双人双锁，并建立电子出入库台账，如实记录品名、规格、数量、领用人和领用时间，做到账物相符。"),
        ("article", "第九条  危险化学品领用实行线上申请、导师审批、双人领取、双人使用、双人归还。单次领用量不得超过3个工作日用量；实验结束后剩余试剂应当于当日退回周转库，不得在实验室存放。"),
        ("article", "第十条  使用危险化学品的实验应当两人以上在场，做好实验记录和交接班记录；过夜实验须经实验室负责人审批并安排值守，严禁单人开展危险性实验。"),
        ("article", "第十一条  危险废物应当分类收集、定点存放，粘贴规范标签，通过学校危险废物管理平台登记，执行转移联单制度，由具有相应资质的单位统一处置。严禁随意倾倒、填埋、焚烧或者排入下水道。"),
        ("article", "第十二条  气瓶应当直立固定、分类存放、远离热源，空瓶与实瓶分开，并定期检验；可燃气体与助燃气体不得混放，使用现场应当保持通风并配备泄漏报警装置。"),
        ("chapter", "第四章  安全培训与准入"),
        ("article", "第十三条  新进入实验室的教师、学生和工作人员必须接受不少于12学时的实验室安全培训，经考试合格后方可进入实验室；在岗人员每年应当接受不少于4学时的安全复训。涉及危险化学品、辐射、生物等特殊领域的，还应当接受专项培训。"),
        ("article", "第十四条  实验室实行准入授权管理，门禁权限与培训考核结果挂钩。未取得准入资格的人员不得进入实验室，不得独自开展实验。"),
        ("article", "第十五条  外来人员进入实验室应当预约登记，由接待人员全程陪同，并遵守实验室安全管理规定；校外人员开展实验活动还应当签订安全责任书。"),
        ("chapter", "第五章  安全检查与应急"),
        ("article", "第十六条  实验室每日开展安全自查，学院每周检查一次，学校每月组织抽查。检查记录应当留存不少于3年，发现隐患立即整改。"),
        ("article", "第十七条  安全隐患实行分级管理。一般隐患立行立改；重大隐患实行挂牌督办，整改期间应当采取停用、隔离等管控措施，并明确责任人、整改期限和验收要求。"),
        ("article", "第十八条  发生实验室安全事故，现场人员应当立即采取应急处置措施，并在30分钟内报告实验室与设备管理处和保卫处，不得迟报、漏报、瞒报；造成人员伤亡的，按照学校突发事件应急预案处置并按规定上报。"),
        ("chapter", "第六章  附则"),
        ("article", "第十九条  本办法由实验室与设备管理处负责解释。"),
        ("article", "第二十条  本办法自2025年9月1日起施行，原《XX大学实验室安全管理办法》（校发〔2023〕12号）同时废止。"),
    ],
}

DOC3 = {
    "file": "XX大学科研经费报销实施细则.docx",
    "title": "XX大学科研经费报销实施细则",
    "doc_no": "校发〔2024〕38号",
    "version_label": "2024年版",
    "effective_from": "2024-05-01",
    "effective_to": None,
    "department": "教务处、科研院、财务处",
    "status": "active",
    "blocks": [
        ("chapter", "第一章  总则"),
        ("article", "第一条  为规范学校科研经费报销管理，提高资金使用效益，根据《高等学校科研经费管理办法》和学校财务管理制度，结合学校实际，制定本细则。"),
        ("article", "第二条  本细则适用于学校承担的纵向科研项目、横向科研项目、校级科研项目及人才项目经费的报销管理。"),
        ("article", "第三条  教务处、科研院负责科研项目的归口管理和预算审核，财务处负责报销审核、会计核算和资金支付，学院（研究院）负责原始票据和业务真实性的初审。"),
        ("chapter", "第二章  开支范围与标准"),
        ("article", "第四条  科研经费直接费用包括设备费、材料费、测试化验加工费、燃料动力费、差旅费、会议费、国际合作与交流费、出版/文献/信息传播/知识产权事务费、劳务费、专家咨询费和其他支出。"),
        ("article", "第五条  国内差旅费按照地区和人员类别标准执行，住宿费凭合法票据在标准限额内据实报销，城市间交通费凭票据报销；伙食补助费标准为100元/人·天，市内交通费标准为80元/人·天。"),
        ("article", "第六条  主办或者承办学术会议，会议费标准为400元/人·天，开支范围包括场地租金、资料印刷、会议用餐等；参会人员的差旅费按照差旅费管理规定另行报销。"),
        ("article", "第七条  劳务费用于支付参与项目研究的研究生、博士后、访问学者以及项目聘用人员的劳务报酬。研究生劳务费标准不超过2000元/人·月，项目聘用人员凭劳动合同和完税凭证报销。"),
        ("article", "第八条  专家咨询费标准为：高级专业技术职称人员800元/人·天，其他人员500元/人·天。咨询活动应当提供邀请函、签到表和书面咨询意见。"),
        ("table", {
            "name": "科研经费部分支出报销标准",
            "headers": ["序号", "支出项目", "标准", "主要凭据"],
            "rows": [
                ["1", "国内差旅住宿费（一类地区）", "650元/人·天", "住宿发票、出差审批单"],
                ["2", "国内差旅住宿费（二类地区）", "550元/人·天", "住宿发票、出差审批单"],
                ["3", "伙食补助费", "100元/人·天", "出差审批单"],
                ["4", "会议费", "400元/人·天", "会议通知、签到表、发票"],
                ["5", "专家咨询费（高级职称）", "800元/人·天", "邀请函、签到表、咨询意见"],
                ["6", "研究生劳务费", "不超过2000元/人·月", "劳务协议、考勤记录"],
            ],
        }),
        ("chapter", "第三章  报销流程"),
        ("article", "第九条  报销票据应当合法有效，发票抬头为“XX大学”，纳税人识别号准确完整；电子发票应当通过财务系统查验并防止重复报销，单张金额1000元以上的票据应当附明细清单。"),
        ("article", "第十条  科研经费报销实行分级审批：单笔5万元以下的，由项目负责人签字、学院科研秘书审核；单笔5万元（含）至20万元的，由学院分管科研负责人审批；单笔20万元（含）至50万元的，由科研院审核、财务处复核；单笔50万元（含）以上的，报分管校领导审批。"),
        ("article", "第十一条  费用原则上应当在发生后3个月内报销，最长不得超过项目执行期结束后1个月。跨年度票据报销应当附情况说明，经科研院和财务处同意后办理。"),
        ("article", "第十二条  报销人应当通过财务网上报销系统提交申请，上传票据影像，经审批通过后按预约时间投递纸质票据，财务处审核无误后办理支付。"),
        ("chapter", "第四章  监督与责任"),
        ("article", "第十三条  严禁虚开发票、虚构劳务、拆分报销、套取或者挪用科研经费。项目负责人对报销业务的真实性、合法性承担第一责任。"),
        ("article", "第十四条  科研院、财务处、审计处应当定期开展联合抽查，重点核查大额支出、劳务费和专家咨询费；发现问题的，暂停该项目报销并限期整改。"),
        ("article", "第十五条  违反本细则的，按照学校科研经费管理和教职工处分有关规定处理；涉嫌违法犯罪的，移送有关机关处理。"),
        ("chapter", "第五章  附则"),
        ("article", "第十六条  本细则由科研院、财务处负责解释。"),
        ("article", "第十七条  本细则自2024年5月1日起施行，原《XX大学科研经费报销管理办法》同时废止。"),
    ],
}

DOC4 = {
    "file": "XX大学教职工请假与考勤管理规定.md",
    "title": "XX大学教职工请假与考勤管理规定",
    "doc_no": "校人〔2024〕9号",
    "version_label": "2024年版",
    "effective_from": "2024-09-01",
    "effective_to": None,
    "department": "人事处",
    "status": "active",
    "md": """# XX大学教职工请假与考勤管理规定

**文号：** 校人〔2024〕9号  
**版本：** 2024年版  
**归口部门：** 人事处  
**施行日期：** 2024年9月1日

## 第一章  总则

**第一条**  为规范教职工请假与考勤管理，维护正常教学、科研和管理秩序，保障教职工合法权益，根据国家和学校有关规定，制定本规定。

**第二条**  本规定适用于学校在编教职工、合同制教职工以及经批准在校内工作的其他人员。

**第三条**  人事处负责全校教职工请假与考勤的归口管理，各单位负责本单位教职工日常考勤、请假初审和销假登记。

## 第二章  请假类型与期限

**第四条**  事假。教职工因私事需要请假的，应当提前申请。全年事假累计一般不超过15个工作日，超过部分按照学校绩效工资有关规定处理。

**第五条**  病假。教职工因病需要请假的，应当提供二级甲等以上医院出具的诊断证明。连续病假5个工作日以上的，还应当附病历和检查报告；病假期间工资待遇按照国家和学校有关规定执行。

**第六条**  婚假、产假、陪产假和丧假。依法办理结婚登记的，享受婚假3个工作日；女教职工生育享受产假158天，难产的增加15天；男教职工享受陪产假15天；直系亲属去世的，享受丧假3个工作日。

**第七条**  年休假。教职工累计工作满1年不满10年的，年休假5天；满10年不满20年的，年休假10天；满20年的，年休假15天。年休假原则上在当年内安排，因工作需要未安排的，经人事处同意可以跨年度安排。

## 第三章  审批权限与流程

**第八条**  请假流程为：本人提出申请，填写请假审批表，经科室（系所）负责人审核后，按照审批权限报学院（部门）或者人事处审批，并办理工作交接。紧急情况可以先电话报告，事后3个工作日内补办手续。

**第九条**  请假审批权限按照请假天数确定：

| 请假天数 | 审批人 | 备案单位 |
| --- | --- | --- |
| 3个工作日以内（含） | 科室（系所）负责人 | 学院（部门）办公室 |
| 3个工作日以上至7个工作日（含） | 学院（部门）分管负责人 | 学院（部门）办公室 |
| 7个工作日以上至15个工作日（含） | 学院（部门）主要负责人 | 人事处 |
| 15个工作日以上 | 分管校领导 | 人事处 |
| 中层及以上干部请假 | 分管校领导或者校长 | 党委办公室、校长办公室 |

**第十条**  请假期间应当做好工作交接。承担教学任务的，应当提前办理调课手续；承担科研项目的，应当明确项目临时负责人。

## 第四章  考勤与销假

**第十一条**  教职工应当遵守工作纪律，按照规定进行考勤打卡或者签到。各单位应当于每月5日前将上月考勤情况报送人事处。

**第十二条**  请假期满后，教职工应当在2个工作日内办理销假手续。需要延长假期的，应当在假期届满前提出申请，经批准后方可继续休假。

**第十三条**  未请假或者请假未获批准擅自离岗的，按照旷工处理。连续旷工15个工作日，或者一年内累计旷工30个工作日的，按照国家和学校有关规定处理。

## 第五章  附则

**第十四条**  本规定由人事处负责解释。

**第十五条**  本规定自2024年9月1日起施行，原《XX大学教职工考勤管理办法》同时废止。
""",
}

DOC5 = {
    "file": "XX学院学生奖学金评定办法.docx",
    "title": "XX学院学生奖学金评定办法",
    "doc_no": "学院〔2024〕5号",
    "version_label": "2024年版",
    "effective_from": "2024-10-01",
    "effective_to": None,
    "department": "XX学院学生工作办公室",
    "status": "active",
    "blocks": [
        ("chapter", "第一章  总则"),
        ("article", "第一条  为激励学生勤奋学习、全面发展，规范学院学生奖学金评定工作，根据《XX大学学生奖学金管理办法》，结合学院实际，制定本办法。"),
        ("article", "第二条  本办法适用于XX学院全日制本科生国家奖学金、国家励志奖学金、校级优秀学生奖学金和单项奖学金的评定。"),
        ("article", "第三条  奖学金评定坚持公开、公平、公正的原则，实行个人申请、班级评议、学院评审、学校审定相结合的工作机制。"),
        ("chapter", "第二章  评选条件"),
        ("article", "第四条  申请奖学金应当具备以下基本条件：拥护中国共产党的领导，遵守国家法律法规和学校规章制度；品行端正，诚实守信，评选年度内无违纪处分；学习态度端正，按规定完成学业；体质健康测试合格。"),
        ("article", "第五条  国家奖学金申请者评选年度学业成绩排名应当位于本专业前10%，且无不及格课程；国家励志奖学金面向家庭经济困难学生，学业成绩排名可以放宽至本专业前30%。"),
        ("article", "第六条  校级优秀学生奖学金一等奖申请者学业成绩排名应当位于本专业前15%，二等奖位于前30%；单项奖学金用于奖励在社会实践、志愿服务、科技创新、文艺体育等方面表现突出的学生。"),
        ("article", "第七条  同一学年内，国家奖学金与国家励志奖学金不可兼得；校级优秀学生奖学金一等奖与二等奖不可兼得。获得国家奖学金或者国家励志奖学金的，可以同时获得单项奖学金。"),
        ("article", "第八条  在科学研究、学科竞赛、创新创业等方面取得突出成绩的，经学院评审委员会三分之二以上成员同意，可以适当放宽学业成绩排名要求，但不得低于本专业前50%。"),
        ("chapter", "第三章  评定程序"),
        ("article", "第九条  奖学金评定工作每年9月启动。学生本人提出申请，填写奖学金申请表，并提交成绩单、获奖证书、志愿服务记录等佐证材料。"),
        ("article", "第十条  班级成立民主评议小组，对申请人的思想品德、学业表现和日常表现进行评议，评议结果在班级范围内公示3个工作日。"),
        ("article", "第十一条  学院成立学生资助评审委员会，对班级推荐人选进行评审，确定拟获奖名单，并在学院范围内公示5个工作日。公示无异议后报学生工作处复核。"),
        ("article", "第十二条  学校学生资助工作领导小组审定获奖名单，发文公布并发放奖学金。学生对评定结果有异议的，可以在公示期内向学院评审委员会提出书面申诉。"),
        ("chapter", "第四章  名额与奖励"),
        ("article", "第十三条  奖学金名额按照年级、专业学生人数比例分配，适当向家庭经济困难学生倾斜。具体名额和奖励标准按照下表执行。"),
        ("table", {
            "name": "XX学院学生奖学金名额与奖励标准",
            "headers": ["奖项", "名额", "金额（元/人·年）", "备注"],
            "rows": [
                ["国家奖学金", "2名", "8000", "按国家下达指标执行"],
                ["国家励志奖学金", "8名", "5000", "面向家庭经济困难学生"],
                ["校级优秀学生奖学金一等奖", "10名", "3000", "学业排名前15%"],
                ["校级优秀学生奖学金二等奖", "20名", "2000", "学业排名前30%"],
                ["单项奖学金", "15名", "1000", "按类别评定"],
            ],
        }),
        ("article", "第十四条  奖学金一次性发放至学生本人银行卡，颁发荣誉证书，获奖情况记入学生学籍档案。"),
        ("chapter", "第五章  附则"),
        ("article", "第十五条  本办法由XX学院学生工作办公室负责解释。"),
        ("article", "第十六条  本办法自2024年10月1日起施行。"),
        ("attachment", "附件1  XX学院学生奖学金申请表"),
        ("table", {
            "name": "XX学院学生奖学金申请表",
            "headers": ["姓名", "学号", "专业年级", "联系电话"],
            "rows": [
                ["", "", "", ""],
                ["申请奖项", "", "学业排名", ""],
                ["主要事迹", "", "", ""],
                ["班级评议意见", "", "签字：      年  月  日", ""],
                ["学院评审意见", "", "签字：      年  月  日", ""],
            ],
        }),
        ("attachment", "附件2  XX学院学生奖学金名额分配表"),
        ("table", {
            "name": "XX学院学生奖学金名额分配表",
            "headers": ["专业", "国家奖学金", "国家励志奖学金", "校级一等奖", "校级二等奖", "单项奖学金"],
            "rows": [
                ["计算机科学与技术", "1", "3", "4", "8", "6"],
                ["软件工程", "1", "3", "4", "8", "6"],
                ["数据科学与大数据技术", "0", "2", "2", "4", "3"],
                ["合计", "2", "8", "10", "20", "15"],
            ],
        }),
    ],
}

DOCS = [DOC1, DOC2, DOC3, DOC4, DOC5]


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
CLOSING = "，。、；：？！”’）】》"
OPENING = "“‘（【《"


def _wrap(text: str, font, size: float, width: float) -> list[str]:
    lines: list[str] = []
    cur = ""
    for ch in text:
        if ch == "\n":
            lines.append(cur)
            cur = ""
            continue
        trial = cur + ch
        if font.text_length(trial, fontsize=size) <= width:
            cur = trial
        elif ch in CLOSING:
            cur = trial
            lines.append(cur)
            cur = ""
        else:
            if cur:
                lines.append(cur)
            cur = ch
    if cur:
        lines.append(cur)
    return lines


def build_pdf(doc: dict, out: Path) -> int:
    # 先用标准参数排版，页数不足 3 页时逐步加大字号/行距，保证 3—4 页。
    for body_size, line_h, gap in ((12.0, 27.0, 8.0), (13.0, 30.0, 10.0), (14.0, 33.0, 12.0)):
        pdf = fitz.open()
        body = fitz.Font(fontfile=MSYH)
        bold = fitz.Font(fontfile=MSYH_BD)
        width = PAGE_W - MARGIN_X * 2
        page = None
        y = 0.0

        def new_page(first: bool = False):
            nonlocal page, y
            page = pdf.new_page(width=PAGE_W, height=PAGE_H)
            page.insert_font(fontname="body", fontfile=MSYH)
            page.insert_font(fontname="bold", fontfile=MSYH_BD)
            y = 118.0 if first else BODY_TOP

        def need(h: float):
            nonlocal y
            if y + h > BODY_BOTTOM:
                new_page()

        def draw(text: str, fontname: str, font, size: float, x: float, color=(0, 0, 0)):
            page.insert_text((x, y), text, fontname=fontname, fontsize=size,
                             color=color, render_mode=0)

        new_page(first=True)
        # 标题区
        title_lines = _wrap(doc["title"], bold, 20, width)
        for line in title_lines:
            w = bold.text_length(line, fontsize=20)
            draw(line, "bold", bold, 20, (PAGE_W - w) / 2)
            y += 31
        y += 6
        meta = [doc["doc_no"], f'{doc["version_label"]}    归口部门：{doc["department"]}',
                f'施行日期：{doc["effective_from"]}']
        for line in meta:
            w = body.text_length(line, fontsize=10.5)
            draw(line, "body", body, 10.5, (PAGE_W - w) / 2, (0.35, 0.35, 0.35))
            y += 17
        y += 14
        page.draw_line((MARGIN_X, y), (PAGE_W - MARGIN_X, y), color=(0.6, 0.6, 0.6), width=0.8)
        y += 26

        for kind, content in doc["blocks"]:
            if kind == "chapter":
                need(48)
                y += 10
                draw(content, "bold", bold, 14, MARGIN_X)
                y += line_h + 8
            elif kind == "article":
                lines = _wrap(content, body, body_size, width - 24)
                need(line_h * len(lines) + gap)
                for i, line in enumerate(lines):
                    x = MARGIN_X + (24 if i == 0 else 0)
                    draw(line, "body", body, body_size, x)
                    y += line_h
                y += gap
            elif kind == "attachment":
                need(40)
                y += 8
                draw(content, "bold", bold, 13, MARGIN_X)
                y += line_h

        count = len(pdf)
        for i, p in enumerate(pdf):
            p.insert_font(fontname="body", fontfile=MSYH)
            label = f"第 {i + 1} 页 / 共 {count} 页"
            w = body.text_length(label, fontsize=9)
            p.insert_text(((PAGE_W - w) / 2, 812), label, fontname="body", fontsize=9,
                          color=(0.4, 0.4, 0.4))
        if 3 <= count <= 4:
            try:
                pdf.subset_fonts()
            except Exception:
                pass
            pdf.set_metadata({
                "title": doc["title"], "author": "XX大学",
                "subject": doc["doc_no"], "keywords": "制度,管理办法,样例",
                "creator": "PolicyRAG sample builder", "producer": "PyMuPDF",
            })
            pdf.save(out, deflate=True, garbage=3)
            pdf.close()
            return count
        pdf.close()
    raise RuntimeError(f"PDF 页数无法控制在 3—4 页：{out}")


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------
def _font(run, name="宋体", size=12, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def _para(doc, text="", *, name="宋体", size=12, bold=False, align=None,
          before=0, after=0, line=1.5, indent=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if indent is not None:
        pf.first_line_indent = Pt(indent)
    if text:
        _font(p.add_run(text), name, size, bold)
    return p


def _table(doc, spec: dict):
    headers, rows = spec["headers"], spec["rows"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(p.add_run(h), "黑体", 10.5, True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i else WD_ALIGN_PARAGRAPH.LEFT
            _font(p.add_run(str(value)), "宋体", 10.5)
    doc.add_paragraph()


def build_docx(doc: dict, out: Path) -> None:
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(2.8)
    normal = d.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    _para(d, doc["title"], name="黑体", size=18, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, before=12, after=8, line=1.3)
    _para(d, doc["doc_no"], name="宋体", size=11,
          align=WD_ALIGN_PARAGRAPH.CENTER, after=4, line=1.2)
    _para(d, f'{doc["version_label"]}    归口部门：{doc["department"]}',
          name="宋体", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, line=1.2)
    _para(d, f'施行日期：{doc["effective_from"]}', name="宋体", size=10.5,
          align=WD_ALIGN_PARAGRAPH.CENTER, after=18, line=1.2)

    for kind, content in doc["blocks"]:
        if kind == "chapter":
            _para(d, content, name="黑体", size=14, bold=True, before=10, after=4, line=1.5)
        elif kind == "article":
            _para(d, content, name="宋体", size=12, before=0, after=4, line=1.6,
                  indent=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif kind == "table":
            _para(d, content["name"], name="黑体", size=10.5, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=3, line=1.2)
            _table(d, content)
        elif kind == "attachment":
            d.add_page_break()
            _para(d, content, name="黑体", size=14, bold=True, before=0, after=8, line=1.4)

    cp = d.core_properties
    cp.title = doc["title"]
    cp.subject = doc["doc_no"]
    cp.author = "XX大学"
    cp.keywords = "制度;样例;PolicyRAG"
    cp.comments = f'{doc["version_label"]}，{doc["department"]}'
    d.save(out)


# --------------------------------------------------------------------------
# 生成与校验
# --------------------------------------------------------------------------
def count_chars(text: str) -> int:
    return len(re.findall(r"[\u4e00-\u9fff]", text))


def text_of(doc: dict) -> str:
    if "md" in doc:
        return doc["md"]
    parts = []
    for kind, content in doc["blocks"]:
        if kind == "table":
            parts.extend(content["headers"])
            for row in content["rows"]:
                parts.extend(str(v) for v in row)
        else:
            parts.append(content)
    return "\n".join(parts)


def main() -> None:
    SAMPLES.mkdir(parents=True, exist_ok=True)
    registry = []
    for doc in DOCS:
        out = SAMPLES / doc["file"]
        if out.suffix.lower() == ".pdf":
            pages = build_pdf(doc, out)
        elif out.suffix.lower() == ".docx":
            build_docx(doc, out)
            pages = 1
        elif out.suffix.lower() == ".md":
            out.write_text(doc["md"], encoding="utf-8")
            pages = 1
        else:
            raise ValueError(out)
        data = out.read_bytes()
        sha = hashlib.sha256(data).hexdigest()
        entry = {k: doc.get(k) for k in (
            "file", "title", "doc_no", "version_label", "effective_from",
            "effective_to", "department", "status")}
        entry["doc_id"] = "doc-" + sha[:12]
        entry["format"] = out.suffix.lower().lstrip(".")
        entry["pages"] = pages
        entry["sha256"] = sha
        registry.append(entry)
        print(f"[ok] {out.name} chars={count_chars(text_of(doc))} pages={pages} size={len(data)}")

    id_by_file = {e["file"]: e["doc_id"] for e in registry}
    old_id = id_by_file[DOC1["file"]]
    for entry in registry:
        entry["supersedes"] = old_id if entry["file"] == DOC2["file"] else None

    meta = {"schema": 1, "updated_at": "2025-09-01", "documents": registry}
    meta_path = SAMPLES / "metadata.json"
    meta_path.write_text(json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"[ok] metadata.json documents={len(registry)}")

    # 回读校验：PDF 文本可提取且含条款；DOCX 可解析且含表格；MD 可读。
    for entry in registry:
        path = SAMPLES / entry["file"]
        if entry["format"] == "pdf":
            with fitz.open(path) as pdf:
                text = "".join(p.get_text() for p in pdf).replace(" ", " ")
                assert entry["pages"] == pdf.page_count
                assert "第 1 页" in text and "第十条" in text, "PDF 文本提取异常"
                print(f"[verify] pdf pages={pdf.page_count} extracted_chars={count_chars(text)}")
        elif entry["format"] == "docx":
            d = Document(path)
            chars = count_chars("\n".join(p.text for p in d.paragraphs))
            print(f"[verify] docx paragraphs={len(d.paragraphs)} tables={len(d.tables)} chars={chars}")
        else:
            chars = count_chars(path.read_text(encoding="utf-8"))
            print(f"[verify] md chars={chars}")


if __name__ == "__main__":
    main()
